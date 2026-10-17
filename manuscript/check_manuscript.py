import re, math, glob, json, sys
import numpy as np, pandas as pd
from docx import Document
import os, subprocess
SH = os.environ.get("RESULTS_DIR") or ("../outputs/tables" if os.path.isdir("../outputs/tables") else "share"); M = "\u2212"
DOCX = sys.argv[1] if len(sys.argv) > 1 else "SOFA2_vs_SOFA1_MIMIC-IV_manuscript.docx"
doc = Document(DOCX)
body = "\n".join(p.text for p in doc.paragraphs)
T = [[[c.text.strip() for c in r.cells] for r in t.rows] for t in doc.tables]
rd = lambda n: pd.read_csv(f"{SH}/{n}.csv")
H, I = rd("auc_results_hospital_death_strict"), rd("auc_results_icu_death_strict_strict")
CH = rd("calib_results_hospital_death_strict"); DD = rd("domain_decomposition_hospital_death_strict"); SP = rd("spo2_vs_pao2_by_group"); MS = rd("audit_missingness_by_subgroup")
CCs, CHs, S2 = "sofa1_community", "sofa1_ett_charted", "sofa2"
SUB = {"sex": ["Male", "Female"], "race": ["White", "Black", "Hispanic/Latino", "Asian", "Other"], "insurance": ["Medicare", "Medicaid", "Private"]}
REF = {"sex": "Male", "race": "White", "insurance": "Medicare"}
fails, passes = [], 0
def f3(x, s=False, d=3):
    r = round(x, d)
    if r == 0: return f"{0:.{d}f}"
    t = f"{abs(r):.{d}f}"; return (M if r < 0 else ("+" if s else "")) + t
def cis(e, l, h, s=False, d=3): return f"{f3(e,s,d)} ({f3(l,s,d)} to {f3(h,s,d)})"
def chk(tag, got, exp):
    global passes
    if got == exp: passes += 1
    else: fails.append((tag, got, exp))
def row(df, fam, p, l=None, comp=None, sc=None):
    m = (df.family == fam) & (df.partition == p)
    if l is not None: m &= df.level == l
    if comp is not None: m &= df.comparator == comp
    if sc is not None: m &= df.score == sc
    assert m.sum() == 1, (fam, p, l, comp, sc, m.sum()); return df[m].iloc[0]
auc = lambda df, s, p="overall", l="All stays": row(df, "auroc", p, l, sc=s)
# ---------------- independent recomputation checks on the CSVs themselves
mx = 0
for df, nm in [(H, "hosp"), (I, "icu")]:
    for p, lv in SUB.items():
        for l in lv:
            for c in (CCs, CHs):
                d_ = auc(df, S2, p, l).estimate - auc(df, c, p, l).estimate
                mx = max(mx, abs(d_ - row(df, "delta_auroc", p, l, c).estimate))
                if l != REF[p]:
                    g = lambda s: auc(df, s, p, REF[p]).estimate - auc(df, s, p, l).estimate
                    mx = max(mx, abs(g(s := S2) - row(df, "gap", p, l, sc=S2).estimate), abs(g(c) - row(df, "gap", p, l, sc=c).estimate))
                    mx = max(mx, abs((g(S2) - g(c)) - row(df, "delta_gap", p, l, c).estimate))
print(f"[A1] max |recomputed - CSV| over delta_auroc, gap, delta_gap (hosp+ICU): {mx:.2e} ->", "PASS" if mx < 1e-9 else "FAIL")
mo = 0
for p, lv in SUB.items():
    for l in lv:
        if l == REF[p]: continue
        for c in (CCs, CHs):
            dsp = lambda s: math.log(row(CH, "oe", p, l, sc=s).estimate) - math.log(row(CH, "oe", p, REF[p], sc=s).estimate)
            mo = max(mo, abs((dsp(S2) - dsp(c)) - row(CH, "delta_calib_disparity_logOE", p, l, c).estimate))
print(f"[A1] max |recomputed - CSV| for calibration delta disparity: {mo:.2e} ->", "PASS" if mo < 1e-9 else "FAIL")
# ---------------- Table 1
N = int(auc(H, S2).n); ND = int(auc(H, S2).deaths)
for p in SUB: pass
sums = {p: (sum(int(auc(H, S2, p, l).n) for l in {"sex": ["Male", "Female"], "race": SUB["race"] + ["Unknown/Declined"], "insurance": SUB["insurance"] + ["Other/Missing"]}[p]),
            sum(int(auc(H, S2, p, l).deaths) for l in {"sex": ["Male", "Female"], "race": SUB["race"] + ["Unknown/Declined"], "insurance": SUB["insurance"] + ["Other/Missing"]}[p])) for p in SUB}
print("[A1] Table 1 totals: N =", N, "deaths =", ND, "| partition sums (n, deaths):", sums, "->", "PASS" if all(v == (N, ND) for v in sums.values()) else "FAIL")
t1 = {r[0].strip(): r for r in T[0]}
def pc(a, b): return f"{100*a/b:.1f}%"
for p, lv in {"sex": ["Male", "Female"], "race": SUB["race"] + ["Unknown/Declined"], "insurance": SUB["insurance"] + ["Other/Missing"]}.items():
    for l in lv:
        r = auc(H, S2, p, l); n, d = int(r.n), int(r.deaths)
        chk(f"T1 {l}", t1[l][1:], [f"{n:,} ({pc(n,N)})", f"{d:,} ({pc(d,n)})"])
cu = {"Medical ICU": "MICU", "Medical/Surgical ICU": "MICUSICU", "Cardiac Vascular ICU": "CVICU", "Surgical ICU": "SICU", "Coronary Care Unit": "CCU", "Trauma SICU": "TSICU"}
for lab, k in cu.items():
    r = auc(rd(f"auc_results_hospital_death_strict_cu-{k}"), S2); n, d = int(r.n), int(r.deaths)
    chk(f"T1 {lab}", t1[lab][1:], [f"{n:,} ({pc(n,N)})", f"{d:,} ({pc(d,n)})"])
# ---------------- Table 2
for i, (lab, df) in enumerate([("In-hospital mortality (primary)", H), ("ICU mortality (secondary)", I)], start=1):
    e = lambda s: auc(df, s); dl = row(df, "delta_auroc_delong", "overall", "All stays", CHs)
    exp = [lab] + [cis(e(s).estimate, e(s).lo, e(s).hi) for s in (CCs, CHs, S2)] + [cis(*[row(df, "delta_auroc", "overall", "All stays", c)[k] for k in ("estimate", "lo", "hi")], True) for c in (CCs, CHs)] + ["<0.001" if dl.p_boot < 0.001 else f"{dl.p_boot:.3f}"]
    chk(f"T2 row{i}", T[1][i], exp)
# ---------------- Table 3 / 4 / 5
t3 = {r[0].strip(): r for r in T[2]}; t4 = {r[0]: r for r in T[3]}; t5 = {r[0].strip(): r for r in T[4]}
for p, lv in SUB.items():
    for l in lv:
        r = auc(H, S2, p, l); lab = l + ("*" if l == REF[p] else "")
        exp = [f"{int(r.n):,} ({int(r.deaths):,})", f3(auc(H, CCs, p, l).estimate), f3(auc(H, CHs, p, l).estimate), f3(r.estimate)] + [cis(*[row(H, "delta_auroc", p, l, c)[k] for k in ("estimate", "lo", "hi")], True) for c in (CCs, CHs)]
        chk(f"T3 {lab}", t3[lab][1:], exp)
        e5 = [cis(*[row(CH, "oe", p, l, sc=s)[k] for k in ("estimate", "lo", "hi")], False, 2) for s in (CCs, CHs, S2)]
        e5 += ["reference", "reference"] if l == REF[p] else [cis(*[row(CH, "delta_calib_disparity_logOE", p, l, c)[k] for k in ("estimate", "lo", "hi")], True) for c in (CCs, CHs)]
        chk(f"T5 {lab}", t5[lab][1:], e5)
        if l != REF[p]:
            g = {s: auc(H, s, p, REF[p]).estimate - auc(H, s, p, l).estimate for s in (CCs, CHs, S2)}
            e4 = [f3(g[CCs], True), f3(g[CHs], True), f3(g[S2], True)] + [cis(*[row(H, "delta_gap", p, l, c)[k] for k in ("estimate", "lo", "hi")], True) for c in (CCs, CHs)] + [f"{f3(abs(g[S2])-abs(g[CHs]),True)} / {f3(abs(g[S2])-abs(g[CCs]),True)}"]
            chk(f"T4 {REF[p]}-{l}", t4[f"{REF[p]} \u2212 {l}"][1:], e4)
# ---------------- supplement tables
sec = lambda f: rd("auc_results_hospital_death_lookback6" if f == "lookback6" else f"auc_results_hospital_death_strict_{f}")
var = [("Headline (broad sedation, strict window)", H), ("Infusion-only sedation", sec("s2-strict")), ("6-hour pre-admission lookback", sec("lookback6")), ("Norepinephrine dose halved", sec("s2-strict_sed_broad_ne_half")), ("Without criteria-based RRT rule", sec("sofa2_total_norrt_criteria")), ("Complete-case stays", sec("cc")), ("Age under 65", sec("age-lt65"))]
for k, (lab, df) in enumerate(var, start=1):
    o = auc(df, S2); d_ = row(df, "delta_auroc", "overall", "All stays", CHs)
    exp = [f"{int(o.n):,} ({int(o.deaths):,})", f3(o.estimate), cis(d_.estimate, d_.lo, d_.hi, True)] + [f3(row(df, "delta_gap", p, l, CHs).estimate, True) for p, l in [("sex", "Female"), ("race", "Black"), ("insurance", "Private")]]
    chk(f"S1 {lab}", T[5][k][1:], exp)
for k, (lab, key) in enumerate(cu.items(), start=1):
    df = rd(f"auc_results_hospital_death_strict_cu-{key}"); o = auc(df, S2)
    d_ = row(df, "delta_auroc", "overall", "All stays", CHs); gs, gb = row(df, "delta_gap", "sex", "Female", CHs), row(df, "delta_gap", "race", "Black", CHs)
    exp = [f"{int(o.n):,} ({int(o.deaths):,})", f"{f3(auc(df,CCs).estimate)} / {f3(auc(df,CHs).estimate)} / {f3(o.estimate)}", cis(d_.estimate, d_.lo, d_.hi, True), cis(gs.estimate, gs.lo, gs.hi, True), cis(gb.estimate, gb.lo, gb.hi, True)]
    chk(f"S2 {lab}", T[6][k][1:], exp)
s3 = [r for r in T[7]]
for blk, comp in [(1, CHs), (None, CCs)]:
    for r in s3:
        pass
doms = ["brain", "respiratory", "respiratory (P/F only)", "cardiovascular", "liver", "kidney", "hemostasis"]
cur = None; cnt = 0
for r in s3[1:]:
    if r[0].startswith("Comparator"): cur = CHs if "as charted" in r[0] else CCs; continue
    d = r[0].strip(); g = lambda fam, p, l: DD[(DD.sofa1_comparator == cur) & (DD.swapped_domain == d) & (DD.family == fam) & (DD.partition == p) & ((DD.level == l) if l else True)].iloc[0]
    a, b, c, e = g("delta_auroc", "overall", "All stays"), g("delta_gap", "sex", "Female"), g("delta_gap", "race", "Black"), g("delta_gap", "insurance", "Private")
    chk(f"S3 {cur}/{d}", r[1:], [cis(x.estimate, x.lo, x.hi, True) for x in (a, b, c, e)]); cnt += 1
for k, r_ in enumerate(SP.itertuples(), start=1):
    exp = [f"{int(r_.n_stays_in_group):,}", f"{r_.pct_resp_from_spo2_only:.1f}%", f"{int(r_.n_stays_paired):,}", f"{int(r_.n_pairs):,}", f"{r_.mean_diff:.2f} ({r_.lo:.2f} to {r_.hi:.2f})", f"{r_.pct_sf_lower:.1f}% / {r_.pct_sf_higher:.1f}%"]
    chk(f"S4 {r_.partition}/{r_.level}", T[8][k][1:], exp)
for k, r_ in enumerate(MS.itertuples(), start=1):
    chk(f"S5 {r_.group}/{r_.level}", T[9][k][1:], [f"{int(r_.n):,}", f"{r_.s1_resp:.1f} / {r_.s2_resp:.1f}", f"{r_.s1_liver:.1f} / {r_.s2_liver:.1f}", f"{r_.s1_complete:.1f} / {r_.s2_complete:.1f}"])
k = 0
for p, lv in SUB.items():
    for l in lv:
        k += 1; r = auc(I, S2, p, l)
        exp = [f"{int(r.n):,} ({int(r.deaths):,})", f3(r.estimate)] + [cis(*[row(I, "delta_auroc", p, l, c)[x] for x in ("estimate", "lo", "hi")], True) for c in (CHs, CCs)]
        exp += ["reference", "reference"] if l == REF[p] else [cis(*[row(I, "delta_gap", p, l, c)[x] for x in ("estimate", "lo", "hi")], True) for c in (CHs, CCs)]
        chk(f"S6 {p}/{l}", T[10][k][1:], exp)
print(f"[A1] table cells compared against independent recomputation: {passes} rows PASS, {len(fails)} FAIL")
for f in fails[:20]: print("   FAIL", f)
# ---------------- text claims (abstract / results)
def has(s): return s in body
tok = {"N": f"{N:,}", "patients": "65,354", "deaths": f"{ND:,}", "11.1%": "11.1%", "ICU deaths 6,226 (7.3%)": f"{int(auc(I,S2).deaths):,} (7.3%)"}
for k, v in tok.items(): print(f"[A1] text contains {k}: {v} ->", "PASS" if has(v) else "FAIL")
for s in (CCs, CHs, S2): print(f"[A1] AUROC {s} {f3(auc(H,s).estimate)} in text ->", "PASS" if has(f3(auc(H, s).estimate)) else "FAIL")
dh = {c: row(H, "delta_auroc", "overall", "All stays", c) for c in (CCs, CHs)}
for c, r in dh.items(): print(f"[A1] abstract/text dAUROC vs {c} {f3(r.estimate,True)} ({f3(r.lo,True)} to {f3(r.hi,True)}) ->", "PASS" if has(f"{f3(r.estimate,True)} ({f3(r.lo,True)} to {f3(r.hi,True)})") else "FAIL")
# decimal tokens in body text that are not found in any CSV value (rounded 2/3 dp, signed or not) -> manual review list
pool = set()
for fpath in glob.glob(f"{SH}/*.csv"):
    df = pd.read_csv(fpath)
    for col in df.select_dtypes("number").columns:
        for v in df[col].dropna().unique():
            for d in (1, 2, 3): pool.add(f3(float(v), False, d).replace(M, "")); pool.add(f"{abs(float(v)):.{d}f}")
txt = re.sub(r"\[\d(,\d)*\]|commit[s]? [0-9a-f, and]+|[0-9a-f]{7}", "", body)
unm = sorted({t.lstrip("+" + M) for t in re.findall(r"[+\u2212-]?\d+\.\d+", txt)} - pool)
print("[A1] decimal tokens in text not present in any CSV value (derived numbers need manual review):", unm)

# ---------------- extra checks for the final document
import glob
seq = []
main = body.split("Declarations")[0]
for m_ in re.finditer(r"\[(\d(?:,\d)*)\]", main):
    for n_ in m_.group(1).split(","):
        if int(n_) not in seq: seq.append(int(n_))
chk("citations appear in order of first citation", seq, sorted(seq))
am = [json.load(open(f)) for f in glob.glob(f"{SH}/analysis_manifest_*.json")]; cm = [json.load(open(f)) for f in glob.glob(f"{SH}/calibration_manifest_*.json")]
chk("one B and one seed across analysis manifests", (len({m["B"] for m in am}), len({m["seed"] for m in am})), (1, 1))
chk("text states the analysis replicate count", f"{am[0]['B']:,} replicates" in body, True)
chk("text states the calibration replicate count", f"{cm[0]['B']:,} replicates" in body, True)
coh_path = os.environ.get("COHORT_FILE") or os.path.join(SH, "cohort_PRIVATE_do_not_share.csv")
if os.path.exists(coh_path):
    c_ = pd.read_csv(coh_path); ex = f"{c_['age'].median():.0f} ({c_['age'].quantile(.25):.0f}\u2013{c_['age'].quantile(.75):.0f})"
    chk("Table 1 age row vs cohort file", t1["Age, median (IQR)"][1], ex)
    chk("cohort file rows == analysed stays", len(c_), N)
gl = subprocess.run(["git", "-C", "..", "log", "--format=%h"], capture_output=True, text=True)
if gl.returncode == 0: chk("every commit in Table S7 exists in git log", {r[0] for r in T[11][1:]} <= set(gl.stdout.split()), True)
print(f"[CHECK] total comparisons: {passes + len(fails)}; PASS {passes}; FAIL {len(fails)}")
for f_ in fails: print("   FAIL", f_)
print("ALL CHECKS PASSED" if not fails else "CHECKS FAILED - do not use this document")
sys.exit(1 if fails else 0)
